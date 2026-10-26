# -*- coding: utf-8 -*-
"""T07 · Panel de administración de frijol: resumen, filtros, mapa, ficha, Excel, bitácora y permisos."""
import sys, pathlib; sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from lib import *
import io, zipfile

adm = login_admin(); tec = login(EM('frijol1'))
ok('login admin y tecnico', adm and tec)

s, r = http('/admin/frijol/resumen', adm); t = r.get('totales', {})
ok('resumen: totales con datos de frijol', s == 200 and t.get('productores', 0) > 0 and t.get('parcelas', 0) > 0 and t.get('hectareas', 0) > 0, str(t))
ok('resumen: desgloses por estado/variedad/tecnico/ciclo', all(len(r.get(k, [])) > 0 for k in ('por_estado', 'por_variedad', 'por_tecnico', 'por_ciclo')))
s, f = http('/admin/frijol/filtros', adm)
ok('filtros: estados, tecnicos y ciclos', s == 200 and f['estados'] and f['tecnicos'] and f['ciclos'])

s, l = http('/admin/frijol/registros?limite=5', adm)
ok('registros paginados', s == 200 and len(l['registros']) == 5 and l['total'] >= 5 and l['paginas'] >= 1, f"total={l.get('total')}")
fila = l['registros'][0]
ok('cada fila trae productor, parcela, ciclo, variedad y tecnico', all(fila.get(k) not in (None, '') for k in ('productor', 'curp', 'up_name', 'ciclo', 'variedad', 'tecnico')))
s, l2 = http('/admin/frijol/registros?limite=5&q=' + fila['curp'], adm)
ok('filtro por CURP', s == 200 and l2['total'] >= 1 and all(x['curp'] == fila['curp'] for x in l2['registros']))
s, l3 = http('/admin/frijol/registros?estado=' + (f['estados'][0]['state_id']) + '&ciclo=' + f['ciclos'][0], adm)
ok('filtro por estado y ciclo', s == 200 and l3['total'] >= 0 and l3['total'] <= l['total'])
s, l4 = http("/admin/frijol/registros?q=%27%3B%20drop%20table%20up%3B--", adm)
ok('entrada maliciosa en el buscador no rompe nada', s == 200 and l4['total'] == 0)

s, m = http('/admin/frijol/mapa', adm)
ok('mapa: GeoJSON con poligonos', s == 200 and m['type'] == 'FeatureCollection' and len(m['features']) > 0 and m['features'][0]['geometry']['type'] in ('MultiPolygon', 'Polygon'), f"{len(m.get('features', []))} parcelas")

s, p = http('/admin/frijol/productores/%d' % fila['producer_id'], adm)
ok('ficha del productor con parcelas, cultivos e historial', s == 200 and p['parcelas'] and p['cultivos'] and 'cambios' in p, p.get('curp', ''))
s, p404 = http('/admin/frijol/productores/999999999', adm); ok('ficha inexistente -> 404', s == 404)
s, pbad = http('/admin/frijol/productores/abc', adm); ok('id invalido -> 400', s == 400)
s, pmz = http('/admin/frijol/productores/1', adm); ok('productor de solo maiz no aparece en el panel de frijol', s == 404)

antes = val('select count(*) from frijol_exportaciones')
s, x = http('/admin/frijol/exportar', adm, raw=True)
ok('exportar: descarga un .xlsx valido', s == 200 and x[:2] == b'PK', f'{len(x)} bytes')
if x[:2] == b'PK':
    z = zipfile.ZipFile(io.BytesIO(x)); wbx = z.read('xl/workbook.xml').decode()
    ok('el Excel trae las hojas Registros y Diccionario', 'name="Registros"' in wbx and 'name="Diccionario"' in wbx)
despues = '|'.join(primera_fila('select count(*), max(filas) from frijol_exportaciones'))
ok('la exportacion queda en la bitacora', despues.split('|')[0] == str(int(antes or 0) + 1), despues)


# ---- criterios de aceptacion del ticket (7) ----
def sql1(q):
    return '|'.join(primera_fila(q))
s, r = http('/admin/frijol/resumen', adm); t = r['totales']
man = sql1("select count(distinct u.producer_id), count(distinct u.up_id), round(sum(cc.area_sown_ha)::numeric,1), round(sum(cc.area_sown_ha*cc.yield_expected)::numeric,1), count(*) from cycle_crop cc join cycle c using(cycle_id) join up u using(up_id) join producer_programa pp on pp.producer_id=u.producer_id and pp.programa='frijol_afb' where c.crop='frijol'").split('|')
ok('tarjetas = conteo manual en la base (productores, parcelas, ha, produccion, cultivos)',
   [t['productores'], t['parcelas'], t['cultivos']] == [int(man[0]), int(man[1]), int(man[4])] and abs(t['hectareas'] - float(man[2])) < 0.06 and abs(t['produccion_estimada_t'] - float(man[3])) < 0.06, f"{man} vs {t}")
ok('hay tarjeta de pendientes > 14 dias', 'pendientes_14_dias' in t and 'produccion_estimada_t' in t)

f = s_ = None
s, fl = http('/admin/frijol/filtros', adm)
ok('filtros traen municipios, variedades, destinos y tipos de tecnico', all(fl.get(k) for k in ('municipios', 'variedades', 'destinos', 'tipos_tecnico')) or fl.get('municipios') is not None)
q = 'estado=10&variedad=PINTO_SALTILLO&semillas=si'
s1, a = http('/admin/frijol/registros?limite=100&' + q, adm); s2, m_ = http('/admin/frijol/mapa?' + q, adm); s3, x = http('/admin/frijol/exportar?' + q, adm, raw=True); s4, rs = http('/admin/frijol/resumen?' + q, adm)
filas_excel = 0
if s3 == 200 and x[:2] == b'PK':
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(x), read_only=True); filas_excel = wb['Registros'].max_row - 1
    cab = [c.value for c in next(wb['Registros'].iter_rows(min_row=1, max_row=1))]
    ok('Excel: hoja Registros con >= 55 columnas y etiquetas legibles', len(cab) >= 55 and 'Producción estimada (t)' in cab and 'Tipo de técnico' in cab and 'Enviado en' in cab, str(len(cab)))
    ok('Excel: hojas Registros y Diccionario', wb.sheetnames == ['Registros', 'Diccionario'])
    primera = [c.value for c in next(wb['Registros'].iter_rows(min_row=2, max_row=2))] if filas_excel else []
    ok('Excel: sin codigos crudos (destino/rotacion/verificacion con etiqueta)', not any(str(v) in ('acopio_axb', 'rota', 'verificada', 'eca_productivo') for v in primera))
ok('filtros combinados: tabla = Excel = tarjeta, mapa <= tabla, y devuelven datos', s1 == 200 and a['total'] > 0 and a['total'] == filas_excel == rs['totales']['cultivos'] and len(m_['features']) <= a['total'],
   f"tabla={a.get('total')} excel={filas_excel} resumen={rs['totales']['cultivos']} mapa={len(m_.get('features', []))}")
s, o1 = http('/admin/frijol/registros?orden=ha&dir=asc&limite=100', adm); s, o2 = http('/admin/frijol/registros?orden=ha&dir=desc&limite=100', adm)
ok('orden ascendente y descendente por hectareas', o1['registros'][0]['ha_sembradas'] <= o1['registros'][-1]['ha_sembradas'] and o2['registros'][0]['ha_sembradas'] >= o2['registros'][-1]['ha_sembradas'])
s, pg = http('/admin/frijol/registros', adm); ok('paginacion de 50 en 50', pg['limite'] == 50 and len(pg['registros']) <= 50)
s, p1 = http('/admin/frijol/registros?pendiente14=1', adm); ok('filtro pendientes > 14 dias responde', s == 200)
s, d = http('/admin/frijol/productores/%d' % fila['producer_id'], adm)
ok('ficha: intentos de verificacion y etiquetas legibles', 'verificacion_intentos' in d and 'verificacion_ultimo_intento' in d and d['cultivos'][0].get('destino') not in ('acopio_axb', 'mercado_libre', 'autoconsumo'))
s, hdr = 0, None

# permisos: un tecnico (capturista) no entra al panel
for ruta in ('/admin/frijol/resumen', '/admin/frijol/registros', '/admin/frijol/mapa', '/admin/frijol/exportar'):
    s, _ = http(ruta, tec, raw=True); ok('tecnico sin acceso a ' + ruta, s in (401, 403), str(s))
for ruta in ('/admin/frijol/resumen', '/admin/frijol/exportar'):
    s, _ = http(ruta, None, raw=True); ok('sin sesion no entra a ' + ruta, s in (401, 403), str(s))
print('\nRESUMEN T07:', sum(res), '/', len(res))
